"""Convierte los checklists 2027 (xlsx de Talento) a fixtures YAML revisables (spec 007).

Uso:  python scripts/checklists_2027_to_yaml.py
Lee   docs/Checklists/Versión 2027/*.xlsx  y los títulos v1 de fixtures/questionnaires/*.yaml.
Escribe fixtures/questionnaires/2027/<AREA>.yaml (fuente que carga `import_questionnaire_version`).

Reglas de título (research.md, D3), en orden:
  a) Si la v1 (de cualquier puesto) tiene el mismo texto con un título específico «Categoría — Subtema», se conserva.
  b) Si la hoja trae Categoría (ID Lead): «Categoría — Actividad».
  c) Si la Actividad se repite en la hoja: «Actividad — Subtema» con el subtema curado en SUBTOPICS.
  d) En otro caso, la Actividad tal cual.
Falla si queda un título repetido dentro de un cuestionario o un subtema sin curar.
"""

import difflib
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "Checklists" / "Versión 2027"
V1 = ROOT / "fixtures" / "questionnaires"
OUT = V1 / "2027"

AREAS = {"Científicos": "CD", "Ingenieros": "ID", "Product": "PM", "UXUI": "UXUI"}
LEVELS = {"jr": "JR", "mid": "MID", "sr": "SNR", "lead": "LEAD"}
LEVEL_ORDER = ["JR", "MID", "SNR", "LEAD"]

# Columna del Excel -> campo del reactivo.
COLUMNS = {
    "Descripción": "text",
    "Criterio de aceptación": "acceptance_criteria",
    "Entregable": "deliverable",
    "Formato": "deliverable_format",
    "Frecuencia": "frequency",
    "Expectativas por nivel de seniority": "level_expectation",
    "Comentarios": "guidance",
}

# Subtemas curados para Actividades repetidas (clave = inicio normalizado de la descripción).
SUBTOPICS = {
    # Ciencia de Datos — rutinas
    "asistir al daily": "Daily",
    "asistir al touchpoint o one-on-one": "Touchpoint con DS Lead/Mid",
    "asistir al touchpoint semanal con el ds lead": "Touchpoint con DS Lead",
    "asistir al touchpoint semanal con los leads": "Touchpoint con Leads de DE y Producto",
    "asistir al touchpoint con los leads": "Touchpoint con Leads de DE y Producto",
    "asistir al sprint planning": "Sprint planning",
    "liderar el sprint planning": "Sprint planning",
    "asistir al sprint review": "Sprint review y retro",
    "exponer su desempeño en el sprint review": "Sprint review y retro",
    "preparar y liderar el biweekly": "Biweekly de datos con Negocio",
    "liderar el biweekly": "Biweekly de datos con Negocio",
    "coordinar y ejecutar las ceremonias de actualización": "Actualización del backlog",
    "supervisar y coordinar las ceremonias de actualización": "Actualización del backlog",
    # Ciencia de Datos — Lead
    "realizar seguimiento a los ds sr/mid": "Seguimiento a checklists",
    "asegurar que los proyectos a su cargo cumplen con las actividades de la oficina de pmo": "Cumplimiento PMO y estándares de arquitectura",
    "asistir y participar activamente en los dailys": "Dailys",
    "asistir y participar activamente en los sprint planning": "Planning y review",
    "identificar, validar el material y participar en sesiones críticas": "Sesiones críticas con stakeholders senior",
    "ayudar al equipo a mapear, entender y gestionar": "Mapeo de stakeholders",
    "agendar y ejecutar one-on-ones con stakeholders": "One-on-ones con stakeholders senior",
    "anticipar y coordinar la participación del director": "Participación del director de Analítica",
    "crear y ejecutar un plan de transición": "Plan de transición y staffing",
    "asegurar que los proyectos a su cargo cumplen con las actividades correspondientes a la mesa": "Mesa de Talento",
    "monitorear y dar seguimiento al avance de los planes de capacitación": "Planes de capacitación",
    "monitorear la ejecución y calidad de las sesiones de mentoring": "Sesiones de mentoring",
    "dar seguimiento a la planeación y anticipación de los días de vacaciones": "Vacaciones y ausencias",
    "ser el punto de referencia para resolver dudas": "Dudas de procedimientos internos",
    "planear y ejecutar sesiones de aprendizaje técnico": "Sesiones de aprendizaje técnico",
    "identificar y gestionar la implementación de mejoras": "Mejoras a metodologías y estándares",
    "dar soporte comercial a través de la identificación": "Identificación de oportunidades",
    "dar soporte comercial a través de la construcción": "Construcción de propuestas",
    "dar soporte comercial a través del seguimiento a facturación": "Facturación y temas administrativos",
    "coordinar y tener sesiones presenciales": "Sesiones presenciales",
    # Producto — Jr/Mid/Sr
    "mantener un plan de trabajo a nivel fase": "Plan de trabajo",
    "mantener un backlog o inventario": "Backlog de features y entregables",
    "documentar actividades completadas cada día": "Registro de actividades del equipo",
    "generar reportes de avance semanales": "Reporte semanal de avance",
    "documentar y enviar por escrito": "Comunicación de ajustes al plan",
    "validar, documentar y compartir al resto del equipo los estándares": "Estándares de trabajo",
    "actualizar indicadores de avance": "Indicadores del proyecto",
    "mantener una estructura de canal y archivos": "Estructura de canal y archivos",
    "asegurar dailys recurrentes": "Dailys",
    "ejecutar weekly status reports": "Weekly status reports",
    "asegurar sprint plan agendados": "Sprint planning",
    "asegurar sprint review agendados": "Sprint review",
    "asegurar sprint retrospective agendados": "Sprint retrospective",
    "agendar refinamientos": "Refinamientos",
    "contar con diseños de solución a nivel técnico": "Diseño de solución técnico",
    "contar con diseños de solución a nivel negocio": "Diseño de solución de negocio",
    "generar presentaciones de revisión y validación": "Presentaciones de resultados",
    "preparar reuniones internas y con el cliente": "Preparación de reuniones",
    "agendar sesiones one on one de retroalimentación": "One on ones de retroalimentación",
    "comunicar por mensaje individual o chat grupal": "Comunicación de avance y puntos críticos",
    "enviar pre-reads": "Pre-reads",
    "documentar y enviar minutas con validaciones": "Minutas de acuerdos y definiciones",
    "documentar y enviar minutas con las mejoras": "Minutas de mejoras y ajustes",
}


def norm(s):
    return re.sub(r"\s+", " ", (s or "").strip()).lower()


def clean(value):
    text = "" if value is None else str(value).strip()
    text = "\n".join(line.rstrip() for line in text.splitlines())
    return "" if norm(text) in ("", "n/a", "na", "-") else text


def v1_specific_titles():
    """texto normalizado -> título específico («A — B») de cualquier cuestionario v1."""
    found = {}
    for path in V1.glob("*.yaml"):
        for tpl in yaml.safe_load(path.read_text(encoding="utf-8"))["templates"]:
            for q in tpl["questions"]:
                if " — " in q["title"]:
                    found.setdefault(norm(q.get("text")), q["title"])
    return found


def subtopic(text):
    key = norm(text)
    matches = [v for k, v in SUBTOPICS.items() if key.startswith(k)]
    return max(matches, key=len) if matches else None


def best_v1(text, specific):
    key = norm(text)
    if key in specific:
        return specific[key]
    close = difflib.get_close_matches(key, specific.keys(), n=1, cutoff=0.95)
    return specific[close[0]] if close else None


def read_sheet(ws):
    header = {str(c.value).strip(): c.column_letter for c in ws[1] if isinstance(c.value, str)}
    rows = []
    for r in range(2, ws.max_row + 1):
        activity = clean(ws[f"{header['Actividad']}{r}"].value)
        if not activity:
            continue
        row = {"activity": activity, "category": clean(ws[f"{header['Categoría']}{r}"].value) if "Categoría" in header else ""}
        for col, field in COLUMNS.items():
            row[field] = clean(ws[f"{header[col]}{r}"].value) if col in header else ""
        rows.append(row)
    return rows


def build_questions(rows, specific, where):
    repeated = {t for t, n in Counter(r["activity"] for r in rows).items() if n > 1}
    questions, missing = [], []
    for r in rows:
        title = best_v1(r["text"], specific)
        if not title and r["category"]:
            title = f"{r['category']} — {r['activity']}"
        if not title and r["activity"] in repeated:
            sub = subtopic(r["text"])
            if not sub:
                missing.append(f"{where}: «{r['activity']}» :: {r['text'][:70]}")
                continue
            title = f"{r['activity']} — {sub}"
        title = title or r["activity"]
        q = {"title": title[:200]}
        q.update({f: r[f] for f in COLUMNS.values() if r[f]})
        questions.append(q)
    dupes = [t for t, n in Counter(q["title"] for q in questions).items() if n > 1]
    return questions, missing, [f"{where}: título repetido «{t}»" for t in dupes]


def main():
    specific = v1_specific_titles()
    OUT.mkdir(parents=True, exist_ok=True)
    errors = []
    for path in sorted(SRC.glob("*.xlsx")):
        area = next(code for key, code in AREAS.items() if key in path.name)
        wb = openpyxl.load_workbook(path, data_only=True)
        templates = []
        for ws in wb.worksheets:
            if not ws.title.lower().startswith("checklist"):
                continue
            level = LEVELS[ws.title.split()[-1].lower()]
            questions, missing, dupes = build_questions(read_sheet(ws), specific, f"{area}/{level}")
            errors += missing + dupes
            templates.append({"area": area, "level": level, "questions": questions})
            print(f"{area}/{level}: {len(questions)} reactivos")
        templates.sort(key=lambda t: LEVEL_ORDER.index(t["level"]))
        doc = {"source": f"docs/Checklists/Versión 2027/{path.name}", "templates": templates}
        (OUT / f"{area}.yaml").write_text(
            yaml.safe_dump(doc, allow_unicode=True, sort_keys=False, width=100000), encoding="utf-8"
        )
    if errors:
        print("\n".join(errors), file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
