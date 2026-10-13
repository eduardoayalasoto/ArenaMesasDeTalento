"""Vistas de planes por área, Mi plan y seguimiento de planes (spec 006, US2/US3/US5).

Validar entrada (Form) → servicio (`learning_plans`) → renderizar (parcial htmx o página).
"""

import io

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from apps.access import services as access
from apps.access.decorators import requires
from apps.catalog.models import Area, SeniorityLevel
from apps.core.services import learning_plans as plans

from .forms import MyPlanAddForm, PlanForm, PlanItemForm
from .models import CatalogCourse, CourseKind, PersonalPlanItem, PlanCourse, Tag, TrainingPlan
from .views import _apply_errors, _base_ctx, _error_text, _forbidden, _tag_chips, _tag_ids, _toast

User = get_user_model()


# --- Planes por área (US2) -----------------------------------------------------------------


@login_required
@requires("learn.plans.view")
def plan_list(request):
    """Abre el plan vigente del área del usuario; `?todos=1` (o sin plan) muestra la lista (FR-015)."""
    if request.GET.get("todos") != "1":
        current = plans.current_plan_for(request.user)
        if current is not None:
            return redirect("learning:plan_detail", pk=current.pk)
    visible = list(plans.visible_plans(request.user))
    years = sorted({p.year for p in visible}, reverse=True)
    return render(request, "learning/plan_list.html", _base_ctx(
        request, "plans", page_title="Planes de capacitación · Arena Learn",
        by_year=[(y, [p for p in visible if p.year == y]) for y in years],
        can_manage_plans=plans.can_manage(request.user),
    ))


def _plan_filters(request):
    f = {k: request.GET.get(k, "").strip() for k in ("kind", "level") if request.GET.get(k, "").strip()}
    f["tags"] = _tag_ids(request)
    return f


def _plan_courses_ctx(request, plan):
    f = _plan_filters(request)
    groups = plans.plan_groups(plan, f)
    tags = Tag.objects.filter(is_active=True, courses__plan_items__plan=plan).distinct()
    can_myplan = access.has(request.user, "learn.myplan") and plan.is_published
    return {
        "plan": plan, "groups": groups, "f": f, "filtered": bool(f.get("kind") or f.get("level") or f.get("tags")),
        "tag_chips": _tag_chips(request, tags, f["tags"]),
        "count": sum(len(g["items"]) for g in groups),
        "can_manage_plans": plans.can_manage(request.user),
        "can_myplan": can_myplan,
        "in_my_plan": plans.my_plan_course_ids(request.user, plan.year) if can_myplan else set(),
        "item_form": PlanItemForm(plan=plan) if plans.can_manage(request.user) else None,
        "kinds": CourseKind.choices, "levels": SeniorityLevel.objects.order_by("order"),
    }


@login_required
@requires("learn.plans.view")
def plan_detail(request, pk):
    plan = get_object_or_404(TrainingPlan.objects.select_related("area"), pk=pk)
    if not plans.can_view_plan(request.user, plan):
        raise Http404
    ctx = _plan_courses_ctx(request, plan)
    if (request.headers.get("HX-Request") and request.GET.get("partial") == "courses"
            and not request.headers.get("HX-History-Restore-Request")):
        resp = render(request, "learning/_plan_courses.html", ctx)
        query = request.GET.copy()
        query.pop("partial", None)
        resp["HX-Push-Url"] = request.path + (f"?{query.urlencode()}" if query else "")
        return resp
    return render(request, "learning/plan_detail.html", _base_ctx(
        request, "plans", page_title=plan.title, period=plans.period_label(plan),
        other_plans=plans.visible_plans(request.user).exclude(pk=plan.pk), **ctx,
    ))


@login_required
@requires("learn.plans.manage")
def plan_edit(request, pk=None):
    plan = get_object_or_404(TrainingPlan, pk=pk) if pk else None
    form = PlanForm(request.POST or None, instance=plan, initial=None if plan else {
        "year": timezone.localdate().year, "start_month": 1, "end_month": 12,
    })
    if request.method == "POST" and form.is_valid():
        obj = form.save(commit=False)
        obj.type_notes = form.type_notes()
        try:
            obj = plans.save_plan(obj, request.user)
        except ValidationError as exc:
            _apply_errors(form, exc)
        else:
            messages.success(request, "Guardaste el plan." if plan else "Creaste el plan en borrador.")
            return redirect("learning:plan_detail", pk=obj.pk)
    return render(request, "learning/plan_form.html", _base_ctx(
        request, "plans", page_title="Editar plan" if plan else "Nuevo plan", form=form, plan=plan,
    ))


@login_required
@require_POST
@requires("learn.plans.manage")
def plan_status(request, pk):
    plan = get_object_or_404(TrainingPlan, pk=pk)
    action = request.POST.get("action") or request.GET.get("action", "")
    try:
        plan = plans.set_status(plan, action, request.user)
    except ValidationError as exc:
        messages.error(request, _error_text(exc))
    else:
        messages.success(request, {
            "publish": "Publicaste el plan.", "reopen": "Volviste a publicar el plan.",
            "unpublish": "El plan regresó a borrador.", "archive": "Archivaste el plan.",
        }.get(action, "Listo."))
    return redirect("learning:plan_detail", pk=pk)


def _courses_partial(request, plan, message="", variant="success"):
    resp = render(request, "learning/_plan_courses.html", _plan_courses_ctx(request, plan))
    return _toast(resp, message, variant) if message else resp


@login_required
@require_POST
@requires("learn.plans.manage")
def plan_item_save(request, pk):
    """Agrega un curso al plan (htmx: devuelve el bloque de cursos)."""
    plan = get_object_or_404(TrainingPlan, pk=pk)
    form = PlanItemForm(request.POST, plan=plan)
    if not form.is_valid() or not form.cleaned_data.get("course"):
        return _courses_partial(request, plan, "Elige un curso del catálogo.", "error")
    try:
        plans.save_item(plan, request.user, course=form.cleaned_data["course"],
                        levels=form.cleaned_data["levels"], note=form.cleaned_data["note"])
    except ValidationError as exc:
        return _courses_partial(request, plan, _error_text(exc), "error")
    return _courses_partial(request, plan, "Agregaste el curso al plan.")


@login_required
@require_POST
@requires("learn.plans.manage")
def plan_item_edit(request, pk, iid):
    """`action=save|delete|up|down` sobre un curso del plan (htmx)."""
    plan = get_object_or_404(TrainingPlan, pk=pk)
    item = get_object_or_404(PlanCourse.objects.select_related("course"), pk=iid, plan=plan)
    action = request.POST.get("action", "save")
    if action == "delete":
        plans.delete_item(item, request.user)
        return _courses_partial(request, plan, "Quitaste el curso del plan.")
    if action in ("up", "down"):
        plans.move_item(item, action, request.user)
        return _courses_partial(request, plan)
    form = PlanItemForm(request.POST, plan=plan)
    if not form.is_valid():
        return _courses_partial(request, plan, "Revisa los niveles y la nota.", "error")
    plans.save_item(plan, request.user, course=item.course, item=item,
                    levels=form.cleaned_data["levels"], note=form.cleaned_data["note"])
    return _courses_partial(request, plan, "Guardaste el curso del plan.")


# --- Mi plan (US3) -------------------------------------------------------------------------


def _year(request) -> int:
    value = request.GET.get("anio") or request.POST.get("anio") or ""
    return int(value) if value.isdigit() else timezone.localdate().year


@login_required
@requires("learn.myplan")
def my_plan(request):
    year = _year(request)
    view = plans.my_plan(request.user, year)
    return render(request, "learning/my_plan.html", _base_ctx(
        request, "myplan", page_title="Mi plan · Arena Learn", view=view, year=year, owner=request.user,
        is_me=True, years=_plan_years(request.user, year), states=plans.STATE_LABELS,
        area_plan=plans.current_plan_for(request.user, year),
    ))


def _plan_years(user, year):
    years = set(user.personal_plans.values_list("year", flat=True)) | {year, timezone.localdate().year}
    return sorted(years, reverse=True)


@login_required
@require_POST
@requires("learn.myplan")
def my_plan_add(request):
    form = MyPlanAddForm(request.POST)
    if not form.is_valid():
        return HttpResponse(status=400)
    pc = course = None
    if form.cleaned_data.get("plan_course"):
        pc = get_object_or_404(PlanCourse.objects.select_related("plan", "course"), pk=form.cleaned_data["plan_course"])
    else:
        course = get_object_or_404(CatalogCourse, pk=form.cleaned_data["course"])
    try:
        item = plans.add_to_my_plan(request.user, course=course, plan_course=pc)
    except ValidationError as exc:
        resp = render(request, "learning/_my_plan_button.html", {"course": course or pc.course, "pc": pc, "added": False})
        return _toast(resp, _error_text(exc), "error")
    resp = render(request, "learning/_my_plan_button.html", {"course": item.course, "pc": pc, "added": True})
    return _toast(resp, f"Agregaste «{item.course.name}» a Mi plan {item.personal_plan.year}.")


def _own_item(request, iid):
    item = get_object_or_404(
        PersonalPlanItem.objects.select_related("course", "personal_plan", "source__plan__area"), pk=iid,
    )
    if item.personal_plan.user_id != request.user.pk:
        raise PermissionDenied
    return item


def _row(request, item, message="", variant="success"):
    state = plans.item_states([item], request.user)[item.pk]
    resp = render(request, "learning/_my_plan_row.html", {"row": {"item": item, "state": state}, "is_me": True})
    return _toast(resp, message, variant) if message else resp


@login_required
@require_POST
@requires("learn.myplan")
def my_plan_start(request, iid):
    """`mode=request` abre la solicitud con autorización; `mode=free` lo deja En curso (FR-019)."""
    item = _own_item(request, iid)
    if request.POST.get("mode") == "request":
        return redirect(f"{reverse('learning:request_create')}?catalogo={item.course_id}")
    try:
        plans.start_item(item, request.user)
    except ValidationError as exc:
        return _row(request, item, _error_text(exc), "error")
    if request.headers.get("HX-Request"):
        return _row(request, item, "Marcaste el curso como en curso.")
    return redirect("learning:my_plan")


@login_required
@require_POST
@requires("learn.myplan")
def my_plan_complete(request, iid):
    """Cierra con el flujo existente: la solicitud (si hay) o un registro directo con evidencia."""
    item = _own_item(request, iid)
    state = plans.item_states([item], request.user)[item.pk]
    req = state.request
    if req is not None and req.status == req.Status.AUTORIZADA:
        return redirect("learning:request_complete", pk=req.pk)
    if req is not None and req.status in (req.Status.EN_REVISION, req.Status.REQUIERE_AJUSTES, req.Status.BORRADOR):
        return redirect("learning:request_detail", pk=req.pk)
    return redirect(f"{reverse('learning:direct_create')}?catalogo={item.course_id}")


@login_required
@require_POST
@requires("learn.myplan")
def my_plan_discard(request, iid):
    item = _own_item(request, iid)
    discard = request.POST.get("action") != "restore"
    plans.discard_item(item, request.user, discard)
    return _row(request, item, "Descartaste el curso." if discard else "Reactivaste el curso.")


@login_required
@requires("learn.myplan.others")
def person_plan(request, pk):
    owner = get_object_or_404(User.objects.select_related("area", "level"), pk=pk)
    if not plans.can_view_personal_plan(request.user, owner):
        return _forbidden(request, "No tienes acceso al plan de esta persona.")
    year = _year(request)
    return render(request, "learning/my_plan.html", _base_ctx(
        request, "people", page_title=f"Plan de {owner.full_name}", view=plans.my_plan(owner, year), year=year,
        owner=owner, is_me=owner.pk == request.user.pk, years=_plan_years(owner, year), states=plans.STATE_LABELS,
        area_plan=None,
    ))


# --- Seguimiento de planes (US5) --------------------------------------------------------------


def _tracking_filters(request):
    raw = {k: request.GET.get(k, "").strip() for k in ("area", "persona", "anio") if request.GET.get(k, "").strip()}
    out = {"year": int(raw["anio"]) if raw.get("anio", "").isdigit() else timezone.localdate().year}
    if raw.get("area", "").isdigit():
        out["area"] = int(raw["area"])
    if raw.get("persona", "").isdigit():
        out["person"] = int(raw["persona"])
    return raw, out


@login_required
@requires("learn.plans.tracking")
def plan_tracking(request):
    raw, filters = _tracking_filters(request)
    data = plans.tracking_rows(request.user, filters)
    this_year = timezone.localdate().year
    return render(request, "learning/plan_tracking.html", _base_ctx(
        request, "tracking", page_title="Seguimiento de planes · Arena Learn", data=data, f=raw,
        areas=Area.objects.filter(is_active=True),
        people=plans.tracking_people(request.user).order_by("full_name"),
        years=range(this_year + 1, this_year - 3, -1), states=plans.STATE_LABELS,
        export_query=request.GET.urlencode(),
    ))


@login_required
@requires("learn.plans.tracking")
def plan_tracking_export(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    _, filters = _tracking_filters(request)
    wb = Workbook()
    ws = wb.active
    ws.title = "Planes"
    ws.append(plans.TRACKING_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in plans.tracking_export_rows(request.user, filters):
        ws.append(row)
    for idx, header in enumerate(plans.TRACKING_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(14, len(header) + 4)
    for cell in ws["K"][1:]:
        cell.number_format = "dd/mm/yyyy"
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="arena_learn_planes_{filters["year"]}.xlsx"'
    return resp
